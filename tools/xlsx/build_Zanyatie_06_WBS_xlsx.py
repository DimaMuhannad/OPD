#!/usr/bin/env python3
"""Excel-таблица к занятию 6: WBS проекта с автопроверкой — вход для диаграммы Ганта.

Листы: «Как заполнять», «Команда» (участники и нагрузка), «WBS» (рабочий), «Пример — умный горшок».
Столбцы WBS совпадают с оборотом бланка (zanyatiya/Zanyatie_06_WBS_A4_Blank): код, элемент,
что предъявляем, ответственный, длительность; плюс «Зависит от» — заполняется на следующей паре.

Запуск: python3 tools/xlsx/build_Zanyatie_06_WBS_xlsx.py  ->  zanyatiya/Zanyatie_06_WBS_Tablica.xlsx
"""
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "zanyatiya/Zanyatie_06_WBS_Tablica.xlsx"

# Палитра ГУАП (tools/marp/guap-brand/BRAND.md). Цвета — ARGB с FF: при альфе 00 Excel для телефона
# рисует заливку не тем цветом
BLUE, INK, GREEN, RED = "FF005AAA", "FF002C5F", "FF009A49", "FFE70F47"
LIGHT, HEAD, GREY = "FFE6EFF8", "FFD9E6F3", "FF7F8C9A"
FONT = "Arial"

thin = Side(style="thin", color="FFB3B3B3")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

HEADER, FIRST, LAST = 2, 3, 62  # шапка и строки данных WBS

# (столбец, заголовок, ширина)
COLS = [
    ("A", "Код", 7),
    ("B", "Уровень", 8),
    ("C", "Элемент WBS: результат или пакет работ (существительным)", 40),
    ("D", "Тип", 14),
    ("E", "Что предъявляем: предмет, документ, данные", 36),
    ("F", "Ответственный (один)", 20),
    ("G", "Длительность", 12),
    ("H", "Ед.", 9),
    ("I", "Зависит от (коды) — на следующей паре", 16),
    ("J", "Проверка", 30),
    ("K", "недель", 8),
]

POT = [
    ("1", "Прототип датчика с индикатором", "", "", None, ""),
    ("1.1", "Компоненты подобраны и куплены", "список компонентов с ценами, не дороже 1500 ₽", "аналитик", 1, "недель"),
    ("1.2", "Схема собрана и проверена на столе", "фото схемы, показания на трёх образцах почвы", "технический эксперт", 1, "недель"),
    ("1.3", "Программа: порог влажности включает индикатор", "код программы", "технический эксперт", 1, "недель"),
    ("1.4", "Корпус с индикатором", "корпус, установленный в горшок", "лидер", 1, "недель"),
    ("2", "Испытания", "", "", None, ""),
    ("2.1", "Калибровка датчика по контрольному влагомеру", "таблица соответствия показаний", "технический эксперт", 1, "недель"),
    ("2.2", "Десять циклов полива на трёх растениях", "протокол испытаний", "аналитик", 2, "недель"),
    ("3", "Проверка с пользователями", "", "", None, ""),
    ("3.1", "Три владельца растений попробовали прототип", "записи разговоров, цитаты", "коммуникатор", 1, "недель"),
    ("4", "Отчёт и защита", "", "", None, ""),
    ("4.1", "Письменный отчёт", "отчёт", "аналитик", 1, "недель"),
    ("4.2", "Презентация и питч", "презентация, отрепетированный питч", "коммуникатор", 1, "недель"),
    ("5", "Управление проектом", "", "", None, ""),
    ("5.1", "Паспорт, план, журнал проекта", "документы; журнал — запись каждую неделю", "лидер", "весь проект", ""),
]
POT_TEAM = [("лидер", "лидер"), ("аналитик", "аналитик"), ("технический эксперт", "технический эксперт"),
            ("коммуникатор", "коммуникатор")]


def font(**k):
    return Font(name=FONT, size=k.pop("size", 10), color=k.pop("color", INK), **k)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def wbs_sheet(ws, title, rows=None, team_ref=None):
    ws.sheet_view.showGridLines = False
    for col, _, w in COLS:
        ws.column_dimensions[col].width = w
    ws.column_dimensions["K"].hidden = True

    # Верх — две строки: «Проект» и сводка, затем шапка. Закреплены только они: на планшете высокая
    # закреплённая область уезжает за край экрана (замечание 17.10)
    ws["A1"] = "Проект:"
    ws["A1"].font = font(bold=True)
    ws.merge_cells("B1:E1")
    ws["B1"].border = Border(bottom=Side(style="thin", color=INK))
    ws["B1"].font = font(size=11)
    if rows:
        ws["B1"] = "Умный горшок: датчик влажности с индикатором (пример)"
    a, b, c, d, j, k = (f"${x}${FIRST}:${x}${LAST}" for x in "ABCDJK")
    ws["F1"] = (f'="Результатов: "&COUNTIFS({b},1)&"  ·  пакетов: "&COUNTIF({d},"пакет работ")'
                f'&"  ·  замечаний: "&(COUNTIF({j},"?*")-COUNTIF({j},"✓"))'
                f'&"  ·  отчёт и защита: "&IF(COUNTIFS({b},1,{c},"*отч*")>0,"✓","добавьте")'
                f'&"  ·  управление проектом: "&IF(COUNTIFS({b},1,{c},"*управл*")>0,"✓","добавьте")'
                f'&"  ·  сумма: "&ROUND(SUM({k}),1)&" нед. по очереди (реальный срок покажет Ганта)"')
    ws["F1"].font = font(bold=True, color=BLUE)
    ws["F1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells("F1:J1")
    ws.row_dimensions[1].height = 34
    ws.conditional_formatting.add("F1", FormulaRule(
        formula=[f'OR(ISNUMBER(SEARCH("добавьте",$F$1)),COUNTIF({j},"?*")-COUNTIF({j},"✓")>0)'],
        font=Font(color=RED, bold=True)))

    for i, (col, head, _) in enumerate(COLS):
        cell = ws[f"{col}{HEADER}"]
        cell.value = head
        cell.font = font(bold=True, size=9)
        cell.fill = fill(HEAD)
        cell.alignment = CENTER
        cell.border = BOX
    ws.row_dimensions[HEADER].height = 42

    for r in range(FIRST, LAST + 1):
        ws[f"A{r}"].number_format = "@"  # «1.10» не превратится в число 1,1
        ws[f"B{r}"] = f'=IF(A{r}="","",LEN(A{r})-LEN(SUBSTITUTE(A{r},".",""))+1)'
        ws[f"D{r}"] = (f'=IF(A{r}="","",IF(OR(B{r}=1,COUNTIF({a},A{r}&".*")>0),"результат","пакет работ"))')
        ws[f"J{r}"] = (
            f'=IF(A{r}="","",IF(COUNTIF({a},A{r})>1,"код повторяется",'
            f'IF(C{r}="","нет названия",'
            f'IF(D{r}="результат",IF(COUNTIF({a},A{r}&".*")=0,"разложите на пакеты работ","✓"),'
            f'IF(E{r}="","впишите, что предъявляем",'
            f'IF(F{r}="","нужен один ответственный",'
            f'IF(OR(ISNUMBER(SEARCH(",",F{r})),ISNUMBER(SEARCH(" и ",F{r})),LOWER(F{r})="все"),"ответственный — один человек",'
            f'IF(G{r}="","оцените длительность","✓"))))))))'
        )
        ws[f"K{r}"] = f'=IF(AND(D{r}="пакет работ",ISNUMBER(G{r})),G{r}*IF(LEFT(LOWER(H{r}),1)="д",1/7,1),0)'
        for col, _, _ in COLS:
            cell = ws[f"{col}{r}"]
            cell.border = BOX
            cell.alignment = CENTER if col in "ABDGHI" else WRAP
            cell.font = font(color=GREY if col in "BDK" else INK)
        ws[f"I{r}"].fill = fill("FFF4F6F8")

    if rows:
        for n, (code, name, show, who, dur, unit) in enumerate(rows):
            r = FIRST + n
            ws[f"A{r}"], ws[f"C{r}"], ws[f"E{r}"], ws[f"F{r}"] = code, name, show, who
            if dur is not None:
                ws[f"G{r}"] = dur
            ws[f"H{r}"] = unit

    # Выпадающие списки
    dv_unit = DataValidation(type="list", formula1='"недель,дней"', allow_blank=True, showErrorMessage=True,
                             errorStyle="stop", errorTitle="Единица", error="Выберите из списка: недель или дней.")
    ws.add_data_validation(dv_unit)
    dv_unit.add(f"H{FIRST}:H{LAST}")
    if team_ref:
        dv_who = DataValidation(type="list", formula1=team_ref, allow_blank=True, showErrorMessage=True, errorStyle="warning",
                                error="Этого человека нет на листе «Команда». Ответственный — один человек.",
                                errorTitle="Ответственный")
        ws.add_data_validation(dv_who)
        dv_who.add(f"F{FIRST}:F{LAST}")

    # Подсветка: крупные результаты — полужирным на голубом, замечания — красным, «✓» — зелёным
    data = f"A{FIRST}:J{LAST}"
    ws.conditional_formatting.add(data, FormulaRule(formula=[f'$D{FIRST}="результат"'],
                                                    fill=fill(LIGHT), font=Font(bold=True, color=INK)))
    ws.conditional_formatting.add(f"J{FIRST}:J{LAST}", FormulaRule(
        formula=[f'AND(J{FIRST}<>"",J{FIRST}<>"✓")'], font=Font(color=RED, bold=True)))
    ws.conditional_formatting.add(f"J{FIRST}:J{LAST}", FormulaRule(
        formula=[f'J{FIRST}="✓"'], font=Font(color=GREEN, bold=True)))

    ws.freeze_panes = f"A{FIRST}"
    ws.print_title_rows = f"1:{HEADER}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:J{LAST}"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4


def team_sheet(ws):
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCD", (28, 22, 14, 14)):
        ws.column_dimensions[col].width = w
    ws["A1"] = "Команда и нагрузка"
    ws["A1"].font = font(size=14, bold=True, color=BLUE)
    ws["A2"] = ("Впишите участников. Их имена появятся в выпадающем списке «Ответственный» на листе WBS. "
                "Справа видно, сколько пакетов и недель работы досталось каждому.")
    ws["A2"].font = font(size=9, italic=True, color=GREY)
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 30
    for col, head in zip("ABCD", ("Участник (имя и фамилия)", "Роль в команде (список справа)", "Пакетов", "Недель работы")):
        c = ws[f"{col}4"]
        c.value, c.font, c.fill, c.alignment, c.border = head, font(bold=True, size=9), fill(HEAD), CENTER, BOX
    # Роли — список справа: четыре роли первого занятия и пустые строки для своих ролей команды
    ws.column_dimensions["E"].width = 3
    ws.column_dimensions["F"].width = 30
    c = ws["F4"]
    c.value, c.font, c.fill, c.alignment, c.border = "Роли команды (список)", font(bold=True, size=9), fill(HEAD), CENTER, BOX
    for i, role in enumerate(["лидер", "аналитик", "технический эксперт", "коммуникатор", "", "", "", ""]):
        c = ws.cell(row=5 + i, column=6, value=role or None)
        c.border, c.font, c.alignment = BOX, font(), WRAP
        if not role:
            c.fill = fill("FFFFF9E1")
    ws["F13"] = "Свою роль впишите в пустую строку выше — она появится в выпадающем списке «Роль»."
    ws["F13"].font = font(size=9, italic=True, color=GREY)
    ws["F13"].alignment = WRAP
    ws.row_dimensions[13].height = 36
    dv = DataValidation(type="list", formula1="$F$5:$F$12", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    for r in range(5, 11):
        dv.add(f"B{r}")
        ws[f"C{r}"] = f'=IF(A{r}="","",COUNTIF(WBS!$F${FIRST}:$F${LAST},A{r}))'
        ws[f"D{r}"] = f'=IF(A{r}="","",ROUND(SUMIF(WBS!$F${FIRST}:$F${LAST},A{r},WBS!$K${FIRST}:$K${LAST}),1))'
        for col in "ABCD":
            c = ws[f"{col}{r}"]
            c.border, c.font = BOX, font()
            c.alignment = CENTER if col in "CD" else WRAP
    ws["A12"] = "Если у одного человека намного больше недель, чем у других, перераспределите пакеты сейчас, до диаграммы Ганта."
    ws["A12"].font = font(size=9, italic=True, color=GREY)
    ws["A12"].alignment = WRAP
    ws.merge_cells("A12:D12")
    ws.row_dimensions[12].height = 30


def howto_sheet(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 100
    lines = [
        ("h", "Иерархическая структура работ (WBS) — таблица"),
        ("p", "Таблица повторяет оборот бланка WBS. Можно работать на бумаге, можно здесь — "
              "на следующую пару принесите лист или откройте файл: из пакетов будем строить диаграмму Ганта."),
        ("s", "Как заполнять"),
        ("n", "1. Лист «Команда»: впишите участников и выберите роли. В списке «Роли команды» справа — четыре роли "
              "с первого занятия; если команда придумала свои роли, впишите их в пустые строки этого списка — "
              "они появятся в выпадающем меню."),
        ("n", "2. Лист «WBS», строка «Проект»: название и главный результат из паспорта."),
        ("n", "3. Крупные результаты — коды 1, 2, 3… Возьмите выбранные ветви дерева целей. "
              "Добавьте «Отчёт и защита» и «Управление проектом»: эту работу всё равно придётся сделать."),
        ("n", "4. Пакеты работ — коды 1.1, 1.2, 2.1… Каждый пакет пишите строкой сразу под своим результатом. "
              "Называйте существительным: «протокол испытаний», а не «испытать»."),
        ("n", "5. Для каждого пакета: что предъявляем, один ответственный (из списка), длительность и единица — недель или дней."),
        ("n", "6. Уровень, тип и столбец «Проверка» заполняются сами. Красное — что ещё поправить, зелёная галочка — готово."),
        ("n", "7. «Зависит от» оставьте пустым — заполним на следующей паре."),
        ("s", "Проверьте себя"),
        ("n", "1. Все выбранные ветви дерева целей есть в WBS, а того, что записано в границах, нет?"),
        ("n", "2. Есть отчёт с защитой и управление проектом?"),
        ("n", "3. Каждый элемент — результат, который можно предъявить, а не действие?"),
        ("n", "4. Каждый пакет можно поручить одному человеку?"),
        ("s", "Подсказки"),
        ("n", "Пакет делается одним человеком за срок от нескольких вечеров до двух недель. Больше — дробите, меньше часа — объедините с соседним."),
        ("n", "Оценивайте честно: если раньше никто в команде так не делал, срок больше, чем кажется."),
        ("n", "Дат и порядка работ в WBS нет: WBS отвечает на вопрос «что», календарный план — «когда»."),
        ("n", "Образец — на листе «Пример — умный горшок»."),
    ]
    r = 1
    for kind, text in lines:
        if kind == "s":
            r += 1  # пустая строка перед подзаголовком
        c = ws.cell(row=r, column=2, value=text)
        c.alignment = WRAP
        if kind == "h":
            c.font = font(size=16, bold=True, color=BLUE)
            ws.row_dimensions[r].height = 26
        elif kind == "s":
            c.font = font(size=12, bold=True, color=BLUE)
        else:
            c.font = font(size=11)
            ws.row_dimensions[r].height = 30 if len(text) > 95 else 16
        r += 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    wb = Workbook()
    howto_sheet(wb.active)
    wb.active.title = "Как заполнять"
    team_sheet(wb.create_sheet("Команда"))
    wbs_sheet(wb.create_sheet("WBS"), "WBS проекта", team_ref="Команда!$A$5:$A$10")
    ex = wb.create_sheet("Пример — умный горшок")
    wbs_sheet(ex, "Пример: WBS умного горшка", rows=POT)
    ex.sheet_properties.tabColor = GREY
    wb["WBS"].sheet_properties.tabColor = BLUE
    wb.active = 0
    wb.properties.creator = "ГУАП · Кафедра № 3"
    wb.properties.title = "Занятие 6. WBS проекта — таблица"
    wb.save(OUT)
    print("OK", OUT)


if __name__ == "__main__":
    main()
